"""
Comprehensive document generator for Task P003:
1. RESEARCH_PAPER_DRAFT.docx (Academic format, ~7-8 pages, structured sections, verified figures & tables)
2. COLLEGE_PROJECT_REPORT.docx (Formal capstone report format, ~18-22 pages, structured sections, placeholders, verified figures & tables)
3. Word COM conversion to matching PDFs.
4. PyMuPDF inspection and validation.
"""

import sys
import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

PROJECT_ROOT = Path("u:/MInor-Specialization")
SUBMISSION_DIR = PROJECT_ROOT / "reports" / "submission"
FIGURES_DIR = PROJECT_ROOT / "reports" / "experiments" / "figures"

# Palette
PRIMARY_COLOR = RGBColor(26, 54, 93)     # Deep Navy
SECONDARY_COLOR = RGBColor(43, 108, 176) # Slate Blue
TEXT_COLOR = RGBColor(45, 55, 72)        # Charcoal body
MUTED_COLOR = RGBColor(113, 128, 150)    # Gray
BG_HEADER_HEX = "EDF2F7"                 # Soft Gray-Blue

def set_cell_background(cell, hex_color):
    xml_str = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(xml_str))

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="A0AEC0"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="A0AEC0"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        <w:insideV w:val="none"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def apply_base_styles(doc):
    # Set standard margins: 1 inch all around
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Set normal style font to Calibri / 11pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = TEXT_COLOR
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

def add_callout_box(doc, text, title="ASSUMPTION & PROTOCOL BOUNDARY"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_background(cell, "F7FAFC")
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:left w:val="single" w:sz="18" w:space="0" w:color="2B6CB0"/>
        <w:top w:val="none"/>
        <w:right w:val="none"/>
        <w:bottom w:val="none"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    run_t = p.add_run(f"[{title}]\n")
    run_t.bold = True
    run_t.font.size = Pt(9.5)
    run_t.font.color.rgb = SECONDARY_COLOR
    
    run_body = p.add_run(text)
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = TEXT_COLOR
    
    # spacing after table
    sp_p = doc.add_paragraph()
    sp_p.paragraph_format.space_before = Pt(0)
    sp_p.paragraph_format.space_after = Pt(4)

print("Base layout utilities loaded.")
